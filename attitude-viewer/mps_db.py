"""Read-only access to the legacy mission-planning schedule log (tbl_mps_mission_sch).

TEST/COMPARISON PATH ONLY. This is NOT the current source of truth for mission data —
see secrets.toml's [mission_db] comment: it was superseded by mce_db.py's
TB_Selected_Mission_Schedule for the scheduling-time target coordinate (this table has
no target coordinate at all, only the pass's start/end ground-track points). It's wired
up behind the /old landing page purely so the exact same UI can be checked against this
older table, for testing/comparison — never as the default mission source.

Notably, this table's cam_start_time/cam_duration/mis_start_time* columns are the same
raw values that mce_db-era missions' MissionParameterJson.missionSettings gets built
from, so the real camera-window formula (mce_db.compute_camera_window) applies directly.

SECURITY: read-only. Only ever issue SELECT here — never write to this DB.
"""
from __future__ import annotations

import re

try:
    import tomllib
except ModuleNotFoundError:
    import tomli as tomllib
from datetime import datetime, timezone
from pathlib import Path

import pymysql

PROJECT_ROOT = Path(__file__).resolve().parent.parent
SECRETS_PATH = PROJECT_ROOT / "secrets.toml"
TARGET_CATALOG_PATH = PROJECT_ROOT / "data" / "old_server_data.xlsx"

TABLE = "tbl_mps_mission_sch"

# mis_target for the fixed/calibration-style targets is "<PREFIX>-<grid>-<seq>", where
# <PREFIX> is sometimes the catalog's own tag ("MT-N2W4-01", "CV-N3E2-04" — seen as-is
# in rows from 2024) and sometimes a satellite ID substituted in later ("O1A-N4W4-01").
# The catalog's own Code column always uses "MT-" or "CV-" for the same "<grid>-<seq>"
# suffix — and critically, MT and CV are two DIFFERENT lists that collide on ~25 shared
# suffixes with completely unrelated coordinates (thousands of km apart). Trusting
# mis_target's own prefix isn't reliable once it's been overwritten with a satellite ID,
# so every candidate is instead checked against the row's own raw end_lat/end_lon ground
# track and the nearest one wins — self-verifying regardless of what the stored prefix
# says. Confirmed against tbl_mps_mission_sch's ground-track points across the table's
# full date range (back to 2023): the correct candidate lands within a few km, the wrong
# one (when a collision exists) thousands of km off, so "nearest" is never ambiguous.
_TARGET_CODE_RE = re.compile(r"^[A-Za-z0-9]+-([NS]\d+[EW]\d+-\d+)$")
_CATALOG_PREFIXES = ("MT", "CV")


def _load_target_catalog() -> dict:
    """suffix ("<grid>-<seq>") -> list of (lat, lon, name) candidates, one per catalog
    prefix (MT/CV) that defines that suffix — usually just one, occasionally both."""
    import openpyxl

    catalog: dict[str, list[tuple[float, float, object]]] = {}
    if not TARGET_CATALOG_PATH.exists():
        return catalog
    wb = openpyxl.load_workbook(TARGET_CATALOG_PATH, data_only=True, read_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = ws.iter_rows(min_row=2, values_only=True)
    for code, name, lat, lon in rows:
        code = str(code or "")
        prefix = next((p for p in _CATALOG_PREFIXES if code.startswith(p + "-")), None)
        if prefix is None or lat is None or lon is None:
            continue
        suffix = code[len(prefix) + 1:]
        catalog.setdefault(suffix, []).append((float(lat), float(lon), name))
    wb.close()
    return catalog


def _nearest_candidate(candidates, ref_lat, ref_lon):
    if ref_lat is None:
        return candidates[0]
    return min(candidates, key=lambda c: (c[0] - float(ref_lat)) ** 2 + (c[1] - float(ref_lon)) ** 2)


_TARGET_CATALOG = None


def _target_catalog() -> dict:
    global _TARGET_CATALOG
    if _TARGET_CATALOG is None:
        _TARGET_CATALOG = _load_target_catalog()
    return _TARGET_CATALOG

# DB stores lowercase bus names (obs1a/obs1b), the app uses O1A/O1B everywhere else.
SAT_ID_TO_APP = {"obs1a": "O1A", "obs1b": "O1B"}
SAT_ID_FROM_APP = {v: k for k, v in SAT_ID_TO_APP.items()}


def _load_config() -> dict:
    with open(SECRETS_PATH, "rb") as f:
        secrets = tomllib.load(f)
    return secrets["mission_db"]


def _connect():
    cfg = _load_config()
    return pymysql.connect(
        host=cfg["db_host"],
        port=int(cfg.get("db_port", 3306)),
        user=cfg["db_user"],
        password=cfg["db_password"],
        database=cfg["db_name"],
        connect_timeout=8,
        read_timeout=15,
        cursorclass=pymysql.cursors.DictCursor,
        charset="utf8mb4",
    )


def _unix_to_iso(unix_ts: float | None) -> str | None:
    if unix_ts is None:
        return None
    dt = datetime.fromtimestamp(unix_ts, tz=timezone.utc)
    return dt.isoformat(timespec="milliseconds").replace("+00:00", "Z")


def _iso_to_unix(iso_str: str | None) -> float | None:
    if not iso_str:
        return None
    try:
        return datetime.fromisoformat(iso_str.replace("Z", "+00:00")).timestamp()
    except ValueError:
        return None


def _row_to_mission(row: dict) -> tuple[dict, bool]:
    """Returns (mission dict, has_exact_coord) — the latter tells get_missions()
    whether this row still needs the mce_db cross-reference fallback below."""
    mis_start_time = row.get("mis_start_time")
    event_start_unix = None
    if mis_start_time is not None:
        event_start_unix = float(mis_start_time) + float(row.get("mis_start_time_ms") or 0) / 1000.0

    event_end_unix = None
    duration_sec = row.get("duration_sec")
    if event_start_unix is not None and duration_sec is not None:
        event_end_unix = event_start_unix + float(duration_sec)

    # Real camera ON~OFF window — same 3-stage formula as mce_db.compute_camera_window,
    # applied to this table's raw equivalents of utcSec/utcMsec/camStartMsec/camDurationUsec.
    # Ground-station-contact rows have no camera fields at all (cam_start_time is NULL).
    scan_start = cam_start = cam_end = None
    cam_start_time = row.get("cam_start_time")
    cam_duration = row.get("cam_duration")
    if event_start_unix is not None and cam_start_time is not None and cam_duration is not None:
        scan_start = _unix_to_iso(event_start_unix)
        cam_start_unix = event_start_unix + float(cam_start_time) / 1000.0
        cam_end_unix = cam_start_unix + float(cam_duration) / 1_000_000.0
        cam_start = _unix_to_iso(cam_start_unix)
        cam_end = _unix_to_iso(cam_end_unix)

    sat_id_raw = (row.get("sat_id") or "").lower()
    satellite_id = SAT_ID_TO_APP.get(sat_id_raw, sat_id_raw.upper())

    # mis_target holds the scheduleId-shaped string for imaging passes (e.g.
    # "O1B_04262_GGD") and the ground-station name for contact passes (e.g.
    # "KSAT-TROLL") — either way it's the best available identifier/label.
    target = row.get("mis_target")

    # This table itself has no target coordinate — only the pass's start/end
    # ground-track points (see module docstring). For the fixed/calibration-style
    # targets ("<SAT>-<grid>-<seq>"), data/old_server_data.xlsx gives the real,
    # exact coordinate — use it whenever mis_target matches. Otherwise (ad-hoc
    # GGD-scheduled missions like "O1A_15658_GGD" or a raw place name), fall back
    # to the pass's end ground-track point, which was verified against mce_db's
    # real target coords to land ~48km away on average (vs ~220km for the midpoint).
    start_lat, start_lon = row.get("start_lat"), row.get("start_lon")
    end_lat, end_lon = row.get("end_lat"), row.get("end_lon")
    lat = lon = None
    code_match = _TARGET_CODE_RE.match(target) if target else None
    catalog_candidates = _target_catalog().get(code_match.group(1)) if code_match else None
    catalog_hit = None
    if catalog_candidates:
        catalog_hit = (
            catalog_candidates[0] if len(catalog_candidates) == 1
            else _nearest_candidate(catalog_candidates, end_lat if end_lat is not None else start_lat,
                                     end_lon if end_lat is not None else start_lon)
        )
    catalog_name = None
    if catalog_hit is not None:
        lat, lon = catalog_hit[0], catalog_hit[1]
        catalog_name = catalog_hit[2]
    elif end_lat is not None:
        lat, lon = float(end_lat), float(end_lon)
    elif start_lat is not None:
        lat, lon = float(start_lat), float(start_lon)

    mission = {
        "id": row.get("idx"),
        "scheduleId": target,
        "satelliteId": satellite_id,
        "operationId": None,
        "location": catalog_name or target,
        "aoiId": target,
        "latitude": lat,
        "longitude": lon,
        "eventStart": _unix_to_iso(event_start_unix),
        "eventEnd": _unix_to_iso(event_end_unix),
        "eventStartKST": None,
        "eventEndKST": None,
        "duration": duration_sec,
        "maxEl": None,
        "cloudAmount": None,
        "requestedScanTime": (float(cam_duration) / 1_000_000.0) if cam_duration is not None else None,
        "note": None,
        "missionStatus": row.get("mis_status"),
        "imageStatus": None,
        # Neither exists in this legacy table — no scheduling-time address/polygon,
        # no processed-image result.
        "clientData": None,
        "results": None,
        "scanStart": scan_start,
        "camStart": cam_start,
        "camEnd": cam_end,
    }
    return mission, catalog_hit is not None


def _enrich_from_current_server(missions: list[dict], satellite_id: str, start_iso: str, end_iso: str) -> None:
    """For missions with no catalog-exact coordinate (typically ad-hoc GGD-scheduled
    missions added after the catalog's own coverage window, e.g. after 2026-08-06),
    cross-reference the CURRENT production schedule (mce_db.TB_Selected_Mission_Schedule)
    for the same window and steal its real target lat/lon/address when the same mission
    is found there too — mce_db has been the source of truth since it superseded this
    table, so for any overlap window it's strictly better than the ground-track guess.

    mis_target here often isn't the same string as mce_db's scheduleId (e.g. this table
    may hold a place name like "중국_다롄시_..." while mce_db has "O1A_15642_34" for the
    identical pass), so rows are matched by scanStart instead — confirmed to line up to
    the millisecond between the two tables (this table's mis_start_time+mis_start_time_ms
    against mce_db's own scanStart) since both ultimately derive from the same uplinked
    schedule data.
    """
    import mce_db

    try:
        current_missions = mce_db.get_missions(satellite_id, start_iso, end_iso)
    except Exception:
        return  # best-effort enrichment only — keep the ground-track fallback on failure

    by_scan_start = {}
    for cm in current_missions:
        ts = _iso_to_unix(cm.get("scanStart"))
        if ts is not None:
            by_scan_start[round(ts)] = cm

    for m in missions:
        ts = _iso_to_unix(m.get("scanStart"))
        if ts is None:
            continue
        match = by_scan_start.get(round(ts))
        if not match or match.get("latitude") is None:
            continue
        m["latitude"] = match["latitude"]
        m["longitude"] = match["longitude"]
        addr = ((match.get("clientData") or {}).get("address") or {})
        m["location"] = addr.get("ko") or addr.get("en") or match.get("location") or m["location"]
        m["clientData"] = match.get("clientData")


def get_missions(satellite_id: str, start_iso: str, end_iso: str) -> list[dict]:
    """TEST/비교용 — 구버전 mission planning schedule log(tbl_mps_mission_sch)에서 조회.

    mce_db.get_missions()와 동일한 반환 형태(dict 리스트, 같은 키 구성)라 프론트엔드는
    이 함수가 반환한 데이터도 그대로 렌더링할 수 있다.
    """
    sat_id_raw = SAT_ID_FROM_APP.get(satellite_id.upper(), satellite_id.lower())

    start_dt = datetime.fromisoformat(start_iso.replace("Z", "+00:00"))
    end_dt = datetime.fromisoformat(end_iso.replace("Z", "+00:00"))
    start_epoch = int(start_dt.astimezone(timezone.utc).timestamp())
    end_epoch = int(end_dt.astimezone(timezone.utc).timestamp())

    conn = _connect()
    try:
        with conn.cursor() as cur:
            cur.execute(
                f"""
                SELECT idx, sat_id, mis_start_time, mis_start_time_ms,
                       start_lat, start_lon, end_lat, end_lon,
                       cam_start_time, cam_duration,
                       duration_sec, mis_target, mis_status
                FROM {TABLE}
                WHERE sat_id = %s AND mis_start_time BETWEEN %s AND %s
                      AND cam_start_time IS NOT NULL AND cam_duration IS NOT NULL
                ORDER BY mis_start_time
                """,
                (sat_id_raw, start_epoch, end_epoch),
            )
            rows = cur.fetchall()
    finally:
        conn.close()

    results = [_row_to_mission(r) for r in rows]
    missions = [m for m, _exact in results]
    uncovered = [m for m, exact in results if not exact]
    if uncovered:
        _enrich_from_current_server(uncovered, satellite_id, start_iso, end_iso)
    return missions
